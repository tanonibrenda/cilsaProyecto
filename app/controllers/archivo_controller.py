# archivo_controller.py
import os
from docx import Document
from flask import send_file

def guardar_informacion_en_txt(info):
    nombre_archivo = "archivos/informacion_guardada.txt"
    with open(nombre_archivo, "a") as archivo:
        archivo.write(info + '\n')
    return nombre_archivo

def guardar_informacion_en_word(info):
    nombre_archivo = "archivos/informacion_guardada.docx"
    document = Document()
    document.add_paragraph(info)
    document.save(nombre_archivo)
    return nombre_archivo

def guardar_informacion_en_txt(info):
    nombre_archivo = "informacion_guardada.txt"
    with open(nombre_archivo, "a") as archivo:
        archivo.write(info + '\n')
    return nombre_archivo

def guardar_informacion_en_word(info):
    nombre_archivo = "informacion_guardada.docx"
    document = Document()
    document.add_paragraph(info)
    document.save(nombre_archivo)
    return nombre_archivo

def obtener_archivos_creados():
    archivos = [f for f in os.listdir() if f.endswith("_info.txt") or f.endswith("_info.docx")]
    return archivos


def descargar_archivo(archivo):
    ruta_archivo = os.path.join("archivos", archivo)
    return send_file(ruta_archivo, as_attachment=True, mimetype="application/octet-stream")